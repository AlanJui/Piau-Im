# -*- coding: utf-8 -*-
"""Build a repair-safe SRT calculator. Uses only native Excel LET formulas."""
from pathlib import Path

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

SRC_XLSX = Path(r"C:\Users\AlanJui\work\Piau-Im\tools\SRT計算器.xlsx")
OUT_XLSX = Path(r"C:\Users\AlanJui\work\Piau-Im\tools\SRT計算器-自動計算.xlsx")
SAFE_XLSX = Path(r"C:\Users\AlanJui\work\Piau-Im\tools\SRT計算器-可用.xlsx")
SRT_PATH = Path(r"C:\Users\AlanJui\_tmp\海海人生.srt")

FONT = "Microsoft JhengHei"
THIN = Border(
    left=Side(style="thin", color="D0D5DD"),
    right=Side(style="thin", color="D0D5DD"),
    top=Side(style="thin", color="D0D5DD"),
    bottom=Side(style="thin", color="D0D5DD"),
)
FILL_TITLE = PatternFill("solid", fgColor="1F4E79")
FILL_SECTION = PatternFill("solid", fgColor="2E75B6")
FILL_INPUT = PatternFill("solid", fgColor="FFF2CC")
FILL_CALC = PatternFill("solid", fgColor="C6EFCE")
FILL_NOTE = PatternFill("solid", fgColor="F2F2F2")
FILL_HEADER = PatternFill("solid", fgColor="1F4E79")
FONT_WHITE = Font(name=FONT, size=14, bold=True, color="FFFFFF")
FONT_SECTION = Font(name=FONT, size=11, bold=True, color="FFFFFF")
FONT_LABEL = Font(name=FONT, size=11)
FONT_BOLD = Font(name=FONT, size=11, bold=True)
FONT_SMALL = Font(name=FONT, size=10, color="595959")
FONT_INPUT = Font(name=FONT, size=12, bold=True)
ALIGN_L = Alignment(vertical="center", wrap_text=True)
ALIGN_C = Alignment(horizontal="center", vertical="center")


def srt_time_to_ms(value: str) -> int:
    sign = -1 if value.startswith("-") else 1
    text = value[1:] if value.startswith("-") else value
    hours, minutes, rest = text.split(":")
    seconds, millis = rest.replace(".", ",").split(",")
    total = ((int(hours) * 3600 + int(minutes) * 60 + int(seconds)) * 1000 + int(millis))
    return sign * total


def parse_srt(path: Path):
    raw = path.read_text(encoding="utf-8").strip().replace("\r\n", "\n")
    rows = []
    for block in (b for b in raw.split("\n\n") if b.strip()):
        lines = block.split("\n")
        idx = int(lines[0])
        show, hide = [p.strip() for p in lines[1].split("-->")]
        text = lines[2] if len(lines) > 2 else ""
        trans = lines[3] if len(lines) > 3 else ""
        duration = srt_time_to_ms(hide) - srt_time_to_ms(show)
        rows.append((idx, show, hide, duration, text, trans))
    return rows


def style(cell, fill=None, font=None, align=None, border=True, fmt=None):
    if fill:
        cell.fill = fill
    if font:
        cell.font = font
    cell.alignment = align or ALIGN_L
    if border:
        cell.border = THIN
    if fmt:
        cell.number_format = fmt


def widths(ws, mapping):
    for col, width in mapping.items():
        ws.column_dimensions[col].width = width


def fx_parse_ms(cell: str) -> str:
    # _xlfn.LET is required when writing LET from openpyxl, otherwise Excel shows #NAME?
    return (
        f'_xlfn.LET(s,TEXT({cell},"@"),'
        f"hh,VALUE(LEFT(s,2)),mm,VALUE(MID(s,4,2)),"
        f"ss,VALUE(MID(s,7,2)),ms,VALUE(MID(s,10,3)),"
        f"(hh*3600+mm*60+ss)*1000+ms)"
    )


def fx_fmt_srt(ms_expr: str) -> str:
    return (
        f"_xlfn.LET(n,ROUND({ms_expr},0),sign,IF(n<0,\"-\",\"\"),x,ABS(n),"
        f"hh,INT(x/3600000),rem,x-hh*3600000,mm,INT(rem/60000),"
        f"rem2,rem-mm*60000,ss,INT(rem2/1000),milli,rem2-ss*1000,"
        f"sign&TEXT(hh,\"00\")&\":\"&TEXT(mm,\"00\")&\":\"&TEXT(ss,\"00\")&\",\"&TEXT(milli,\"000\"))"
    )


def fx_show_plus_dur(show_cell: str, dur_cell: str) -> str:
    return "=" + fx_fmt_srt(f"{fx_parse_ms(show_cell)}+{dur_cell}")


def fx_offset_ms(new_show: str, old_show: str) -> str:
    return f"={fx_parse_ms(new_show)}-{fx_parse_ms(old_show)}"


def fx_shift(time_cell: str, no_cell: str, start_cell: str, offset_cell: str) -> str:
    return (
        f"=IF({no_cell}<{start_cell},{time_cell},"
        + fx_fmt_srt(f"{fx_parse_ms(time_cell)}+{offset_cell}")
        + ")"
    )


def build_help_sheet(wb):
    ws = wb.create_sheet("說明", 0)
    ws.sheet_view.showGridLines = False
    ws["A1"] = "SRT 計算器｜使用說明"
    ws.merge_cells("A1:B1")
    style(ws["A1"], fill=FILL_TITLE, font=FONT_WHITE, border=False)
    ws.row_dimensions[1].height = 28

    notes = [
        (3, "立刻能用", "「單句」B 欄綠色格是 Excel 公式，開檔就應算出 Hide／位移，不必先貼 PY。"),
        (4, "黃色儲存格", "只有黃色格要改。時間請用文字：00:00:48,000（逗號後面是毫秒）。"),
        (5, "不要在 B7 按 =PY", "B7／B11／B12 已有公式。若在那裡按 =PY 再 Tab，公式會被清空，變成空的 PY 格。"),
        (6, "若要用 PY()", "到「腳本」複製程式 → 在空白的 D5 或 K2 輸入 =PY 再 Tab → 貼上 → Ctrl+Enter → 輸出選 Excel 值。溢出範圍必須全空，提示改放附註。"),
        (7, "這首歌的數字", "前奏到 00:00:48,000；#2 舊起點 00:00:04,000；位移 44000 毫秒；新 Hide 00:00:52,000。"),
        (8, "批次工作表", "G／H 欄會依單句位移整表重算。#1 歌名預設不位移。"),
        (9, "匯回 Subtitle Edit", "複製批次 I 欄，貼到記事本存成 .srt，再用 File → Open 開啟。"),
    ]
    ws["A2"] = "項目"
    ws["B2"] = "說明"
    style(ws["A2"], fill=FILL_SECTION, font=FONT_SECTION)
    style(ws["B2"], fill=FILL_SECTION, font=FONT_SECTION)
    for row, title, text in notes:
        ws[f"A{row}"] = title
        ws[f"B{row}"] = text
        style(ws[f"A{row}"], fill=FILL_NOTE, font=FONT_BOLD)
        style(ws[f"B{row}"], font=FONT_LABEL)
        ws.row_dimensions[row].height = 36
    widths(ws, {"A": 22, "B": 96})
    ws.freeze_panes = "A3"


def build_single_sheet(wb):
    ws = wb.create_sheet("單句", 1)
    ws.sheet_view.showGridLines = False
    ws["A1"] = "單句：改起點、維持時長"
    ws.merge_cells("A1:B1")
    style(ws["A1"], fill=FILL_TITLE, font=FONT_WHITE, border=False)
    ws.row_dimensions[1].height = 28

    ws["A2"] = "黃色＝輸入。綠色＝自動計算。時間格式必須是 00:00:48,000"
    ws.merge_cells("A2:B2")
    style(ws["A2"], fill=FILL_NOTE, font=FONT_SMALL, border=False)

    ws["A3"] = "項目"
    ws["B3"] = "值"
    style(ws["A3"], fill=FILL_SECTION, font=FONT_SECTION, align=ALIGN_C)
    style(ws["B3"], fill=FILL_SECTION, font=FONT_SECTION, align=ALIGN_C)

    sections = {
        4: "目前",
        8: "改起點、維持時長",
        13: "批次套用",
    }
    labels = {
        5: "目前 Show（#2 舊起點）",
        6: "Duration（毫秒）",
        7: "目前 Hide（自動）",
        9: "新 Show（前奏結束／第一句開口）",
        10: "Duration（毫秒，沿用上面）",
        11: "新 Hide（自動）",
        12: "位移（毫秒）",
        14: "從第幾條開始位移",
        15: "位移（SRT，供核對）",
    }
    for row, title in sections.items():
        ws.merge_cells(f"A{row}:B{row}")
        ws[f"A{row}"] = title
        style(ws[f"A{row}"], fill=FILL_SECTION, font=FONT_SECTION)
        ws.row_dimensions[row].height = 22
    for row, label in labels.items():
        ws[f"A{row}"] = label
        style(ws[f"A{row}"], font=FONT_LABEL)
        style(ws[f"B{row}"])
        ws.row_dimensions[row].height = 24

    ws["B5"] = "00:00:04,000"
    ws["B9"] = "00:00:48,000"
    for addr in ("B5", "B9"):
        style(ws[addr], fill=FILL_INPUT, font=FONT_INPUT, align=ALIGN_C, fmt="@")

    ws["B6"] = 4000
    style(ws["B6"], fill=FILL_INPUT, font=FONT_INPUT, align=ALIGN_C, fmt="#,##0")
    ws["B10"] = "=B6"
    style(ws["B10"], fill=FILL_CALC, font=FONT_INPUT, align=ALIGN_C, fmt="#,##0")
    ws["B14"] = 2
    style(ws["B14"], fill=FILL_INPUT, font=FONT_INPUT, align=ALIGN_C, fmt="#,##0")

    ws["B7"] = fx_show_plus_dur("B5", "B6")
    ws["B11"] = fx_show_plus_dur("B9", "B10")
    ws["B12"] = fx_offset_ms("B9", "B5")
    ws["B15"] = "=" + fx_fmt_srt("B12")
    for addr in ("B7", "B11", "B15"):
        style(ws[addr], fill=FILL_CALC, font=FONT_INPUT, align=ALIGN_C, fmt="@")
    style(ws["B12"], fill=FILL_CALC, font=FONT_INPUT, align=ALIGN_C, fmt="#,##0")

    ws.merge_cells("A17:B17")
    ws["A17"] = "對照 Subtitle Edit"
    style(ws["A17"], fill=FILL_SECTION, font=FONT_SECTION)
    ws.merge_cells("A18:B18")
    ws["A18"] = (
        "在 SE 列表改 #2 的 Show，Hide 不會跟著走，Duration 會變負值。"
        "本表把 Duration 鎖在 B6，只重算 Hide。"
        "預設：舊 00:00:04,000 → 新 00:00:48,000，位移 44000 毫秒，新 Hide 應為 00:00:52,000。"
    )
    style(ws["A18"], fill=FILL_NOTE, font=FONT_LABEL)
    ws.row_dimensions[18].height = 52

    ws["D1"] = "PY() 請貼 D5"
    style(ws["D1"], fill=FILL_TITLE, font=FONT_WHITE, border=False)
    ws["E1"] = "溢出範圍 D5:E9 必須全空"
    style(ws["E1"], fill=FILL_TITLE, font=FONT_WHITE, border=False)
    ws["D2"] = "選 D5 → =PY → Tab → 貼「腳本」B5 → Ctrl+Enter"
    style(ws["D2"], fill=FILL_NOTE, font=FONT_SMALL, border=False)
    ws["E2"] = None
    style(ws["D5"], fill=FILL_CALC, font=FONT_INPUT, border=False)
    ws["D5"].comment = Comment(
        "在此格輸入 =PY 後按 Tab，貼上「腳本」工作表 B5 的程式，再按 Ctrl+Enter。"
        "輸出請改為 Excel 值。"
        "結果會溢出到 D5:E9，這塊範圍請保持空白（含 E5）。",
        "SRT",
    )
    ws["D5"].comment.width = 280
    ws["D5"].comment.height = 110

    widths(ws, {"A": 40, "B": 28, "C": 3, "D": 28, "E": 28})
    ws.freeze_panes = "A4"


def build_batch_sheet(wb, rows):
    ws = wb.create_sheet("批次", 2)
    headers = ["編號", "Show", "Hide", "Duration", "Text", "Translation", "新Show", "新Hide", "可匯出SRT"]
    for col, header in enumerate(headers, 1):
        style(
            ws.cell(1, col, header),
            fill=FILL_HEADER,
            font=Font(name=FONT, size=11, bold=True, color="FFFFFF"),
            align=ALIGN_C,
        )

    last = 1
    for r_i, (idx, show, hide, dur, text, trans) in enumerate(rows, 2):
        last = r_i
        ws.cell(r_i, 1, idx)
        ws.cell(r_i, 2, show).number_format = "@"
        ws.cell(r_i, 3, hide).number_format = "@"
        ws.cell(r_i, 4, dur)
        ws.cell(r_i, 5, text)
        ws.cell(r_i, 6, trans)
        ws.cell(r_i, 7, fx_shift(f"B{r_i}", f"A{r_i}", "單句!$B$14", "單句!$B$12"))
        ws.cell(r_i, 8, fx_shift(f"C{r_i}", f"A{r_i}", "單句!$B$14", "單句!$B$12"))
        ws.cell(r_i, 9, f'=A{r_i}&CHAR(10)&G{r_i}&" --> "&H{r_i}&CHAR(10)&E{r_i}&CHAR(10)&F{r_i}')
        for c in range(1, 10):
            cell = ws.cell(r_i, c)
            cell.font = FONT_LABEL
            cell.border = THIN
            cell.alignment = Alignment(vertical="center", wrap_text=(c == 9))
            if c in (1, 2, 3, 4, 7, 8):
                cell.alignment = ALIGN_C
            if c in (2, 3, 7, 8):
                cell.number_format = "@"
            if c == 4:
                cell.number_format = "#,##0"
            if c in (7, 8):
                cell.fill = FILL_CALC
            if c == 9:
                cell.fill = FILL_NOTE
        ws.row_dimensions[r_i].height = 36

    ws["K1"] = "PY結果（貼K2）"
    style(ws["K1"], fill=FILL_HEADER, font=Font(name=FONT, size=11, bold=True, color="FFFFFF"), align=ALIGN_C)
    ws["L1"] = None
    style(ws["K2"], fill=FILL_CALC, border=False)
    ws["K2"].comment = Comment(
        "在此格輸入 =PY 後按 Tab，貼上「腳本」工作表 B6 的程式，再按 Ctrl+Enter。"
        "輸出請改為 Excel 值。"
        "結果會向右、向下溢出（K2:L 約 62 列），K 與 L 欄此處以下請保持空白。",
        "SRT",
    )
    ws["K2"].comment.width = 280
    ws["K2"].comment.height = 120

    widths(
        ws,
        {"A": 8, "B": 16, "C": 16, "D": 12, "E": 28, "F": 42, "G": 16, "H": 16, "I": 36, "J": 3, "K": 16, "L": 16},
    )
    ws.freeze_panes = "A2"
    ws.row_dimensions[1].height = 22
    return last


PY_SINGLE = '''def srt_to_ms(s):
    text = str(s).strip()
    h, m, rest = text.split(":")
    sec, ms = rest.replace(".", ",").split(",")
    return (int(h) * 3600 + int(m) * 60 + int(sec)) * 1000 + int(ms)

def ms_to_srt(ms):
    ms = int(round(float(ms)))
    sign = "-" if ms < 0 else ""
    ms = abs(ms)
    h, rem = divmod(ms, 3600000)
    mi, rem = divmod(rem, 60000)
    sec, milli = divmod(rem, 1000)
    return f"{sign}{h:02d}:{mi:02d}:{sec:02d},{milli:03d}"

old_show = xl("B5")
dur = xl("B6")
new_show = xl("B9")
new_dur = xl("B10")
offset = srt_to_ms(new_show) - srt_to_ms(old_show)

pd.DataFrame({
    "項目": ["目前Hide", "新Hide", "位移毫秒", "位移SRT"],
    "值": [
        ms_to_srt(srt_to_ms(old_show) + dur),
        ms_to_srt(srt_to_ms(new_show) + new_dur),
        offset,
        ms_to_srt(offset),
    ]
})
'''

PY_BATCH = '''def srt_to_ms(s):
    text = str(s).strip()
    h, m, rest = text.split(":")
    sec, ms = rest.replace(".", ",").split(",")
    return (int(h) * 3600 + int(m) * 60 + int(sec)) * 1000 + int(ms)

def ms_to_srt(ms):
    ms = int(round(float(ms)))
    sign = "-" if ms < 0 else ""
    ms = abs(ms)
    h, rem = divmod(ms, 3600000)
    mi, rem = divmod(rem, 60000)
    sec, milli = divmod(rem, 1000)
    return f"{sign}{h:02d}:{mi:02d}:{sec:02d},{milli:03d}"

df = xl("A1:F62", headers=True)
old_show = xl("單句!B5")
new_show = xl("單句!B9")
start_no = xl("單句!B14")
offset = srt_to_ms(new_show) - srt_to_ms(old_show)

def shift_one(t, no):
    ms = srt_to_ms(t)
    if no >= start_no:
        ms = ms + offset
    return ms_to_srt(ms)

pd.DataFrame({
    "新Show": [shift_one(s, n) for s, n in zip(df["Show"], df["編號"])],
    "新Hide": [shift_one(s, n) for s, n in zip(df["Hide"], df["編號"])],
})
'''


def build_script_sheet(wb):
    ws = wb.create_sheet("腳本", 1)
    ws.sheet_view.showGridLines = False
    ws["A1"] = "複製這裡的 Python，貼進 Excel 的 PY 編輯器"
    ws.merge_cells("A1:B1")
    style(ws["A1"], fill=FILL_TITLE, font=FONT_WHITE, border=False)
    ws.row_dimensions[1].height = 28

    ws["A2"] = "步驟：選 D5 或 K2 → 輸入 =PY → Tab → 貼上右欄腳本 → Ctrl+Enter → 改為 Excel 值。D5:E9 與 K2 右側／下方必須全空。"
    ws.merge_cells("A2:B2")
    style(ws["A2"], fill=FILL_NOTE, font=FONT_SMALL, border=False)
    ws.row_dimensions[2].height = 36

    ws["A3"] = "貼到哪裡"
    ws["B3"] = "Python 腳本（整格複製）"
    style(ws["A3"], fill=FILL_SECTION, font=FONT_SECTION)
    style(ws["B3"], fill=FILL_SECTION, font=FONT_SECTION)

    ws["A4"] = "單句!D5（一次算出 Hide 與位移）"
    ws["B4"] = "整格複製下面 B5"
    style(ws["A4"], font=FONT_BOLD)
    style(ws["B4"], font=FONT_SMALL)
    ws["B5"] = PY_SINGLE.strip()
    style(ws["B5"], fill=FILL_CALC, font=Font(name="Consolas", size=10), align=Alignment(vertical="top", wrap_text=True), border=False)
    ws.row_dimensions[5].height = 240

    ws["A6"] = "批次!K2（溢出 新Show／新Hide）"
    ws["B6"] = PY_BATCH.strip()
    style(ws["A6"], font=FONT_BOLD)
    style(ws["B6"], fill=FILL_CALC, font=Font(name="Consolas", size=10), align=Alignment(vertical="top", wrap_text=True), border=False)
    ws.row_dimensions[6].height = 260

    widths(ws, {"A": 40, "B": 88})
    ws.freeze_panes = "A3"


def main():
    rows = parse_srt(SRT_PATH)
    wb = Workbook()
    default = wb.active
    build_help_sheet(wb)
    build_script_sheet(wb)
    build_single_sheet(wb)
    build_batch_sheet(wb, rows)
    wb.remove(default)
    wb.properties.title = "SRT 計算器"
    wb.properties.creator = "海海人生 MV"
    wb["單句"].sheet_properties.tabColor = "548235"
    wb["批次"].sheet_properties.tabColor = "1F4E79"
    wb["說明"].sheet_properties.tabColor = "BF8F00"
    wb["腳本"].sheet_properties.tabColor = "833C0C"

    script_dir = Path(r"C:\Users\AlanJui\work\Piau-Im\tools\excel_py_scripts")
    script_dir.mkdir(parents=True, exist_ok=True)
    (script_dir / "srt_single.py").write_text(PY_SINGLE.strip() + "\n", encoding="utf-8")
    (script_dir / "srt_batch.py").write_text(PY_BATCH.strip() + "\n", encoding="utf-8")

    saved = []
    for target in (SAFE_XLSX, SRC_XLSX, OUT_XLSX):
        try:
            wb.save(target)
            saved.append(str(target))
        except PermissionError:
            alt = target.with_name(target.stem + "-新.xlsx")
            wb.save(alt)
            saved.append(f"{alt} (原檔使用中)")
    print(f"cues={len(rows)}")
    for item in saved:
        print(item)


if __name__ == "__main__":
    main()
